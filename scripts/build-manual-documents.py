from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "manual-assets"
OUTPUT = DOCS / "generated-manuals"
OUTPUT.mkdir(parents=True, exist_ok=True)

BLUE = "2563EB"
DARK = "172033"
MUTED = "64748B"
LIGHT_BLUE = "EFF6FF"
FONT = "TH Sarabun New"


REQUESTER_IMAGES = {
    "## 0 ภาพรวม Flow การทำงาน": [("program-workflow.png", "แผนภาพที่ 1 Flow การทำงานจริงตาม WorkflowTransition ปัจจุบัน")],
    "*ภาพที่ 1 หน้าแรก*": [("requester-dashboard.png", "ภาพที่ 1 หน้าแรก (Dashboard)")],
    "*ภาพที่ 2 หน้าคำร้องของคุณ*": [("requester-requests.png", "ภาพที่ 2 หน้าคำร้องของคุณ")],
    "*ภาพที่ 3 ข้อมูลทั่วไป*": [("requester-new-step1.png", "ภาพที่ 3 ขั้นตอนข้อมูลทั่วไป")],
    "**• รายละเอียดและไฟล์แนบ**": [("requester-new-step2.png", "ภาพที่ 4 ขั้นตอนรายละเอียดและไฟล์แนบ")],
    "**• ตรวจสอบและส่ง**": [("requester-new-step3.png", "ภาพที่ 5 ขั้นตอนตรวจสอบและส่ง")],
    "#### การติดตามสถานะ": [("requester-detail.png", "ภาพที่ 6 หน้ารายละเอียดและการติดตามสถานะคำร้อง")],
}

APPROVER_IMAGES = {
    "## 0 ภาพรวมเส้นทางของคำร้อง": [("program-workflow.png", "แผนภาพที่ 1 Flow การทำงานจริงตาม WorkflowTransition ปัจจุบัน")],
    '### 1.1 เมนู "งานรออนุมัติ"': [("head-pending-tasks.png", "ภาพที่ 1 รายการที่ต้องอนุมัติหรือดำเนินการ")],
    "### 2.2 การอนุมัติ": [
        ("head-request-detail.png", "ภาพที่ 2 หน้ารายละเอียดคำร้องและปุ่มดำเนินการ"),
        ("head-approve-dialog.png", "ภาพที่ 3 กล่องยืนยันการอนุมัติ"),
    ],
    "### 2.3 การส่งกลับให้แก้ไข": [("head-reject-dialog.png", "ภาพที่ 4 การระบุเหตุผลเพื่อส่งกลับแก้ไข")],
    "### 3.1 การตรวจรอบแรก — สถานะ *รอนำส่งบัญชีตรวจสอบ*": [
        ("accountant-pending-tasks.png", "ภาพที่ 5 รายการของฝ่ายบัญชี ซึ่งแยกงานตรวจรอบแรกและรอบหลังแก้ไข"),
        ("accountant-request-detail.png", "ภาพที่ 6 รายละเอียดคำร้องสำหรับฝ่ายบัญชีตรวจสอบ"),
    ],
    "### 4.1 การอนุมัติ": [
        ("final-pending-tasks.png", "ภาพที่ 7 รายการที่รอผู้อนุมัติสูงสุด"),
        ("final-request-detail.png", "ภาพที่ 8 รายละเอียดคำร้องสำหรับผู้อนุมัติสูงสุด"),
    ],
    "### 4.2 การดูภาพรวม": [("approver-dashboard.png", "ภาพที่ 9 หน้าแดชบอร์ดภาพรวม")],
    "### 4.3 รายงานย้อนหลัง": [("approver-report.png", "ภาพที่ 10 หน้ารายงานย้อนหลัง")],
    "### 5.1 ผู้ดำเนินการแก้ไข — สถานะ *รอ IT ดำเนินการ*": [
        ("it-operator-pending-tasks.png", "ภาพที่ 11 รายการที่รอ IT ดำเนินการ"),
        ("it-operator-request-detail.png", "ภาพที่ 12 รายละเอียดคำร้องสำหรับ IT ดำเนินการ"),
    ],
    "### 5.2 ผู้ปิดงาน — สถานะ *รอ IT ปิดงาน*": [
        ("it-reviewer-pending-tasks.png", "ภาพที่ 13 รายการที่รอ IT ตรวจรับและปิดงาน"),
        ("it-reviewer-request-detail.png", "ภาพที่ 14 รายละเอียดคำร้องสำหรับผู้ตรวจรับงาน IT"),
    ],
    "### 8.2 ลายเซ็น": [("approver-profile.png", "ภาพที่ 15 หน้าโปรไฟล์และการจัดการลายเซ็น")],
}


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("หน้า ")
    set_run_font(run, 11, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_inline(paragraph, text, base_size=16, base_color=DARK):
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`|(?<!\*)\*[^*]+\*(?!\*))")
    for part in filter(None, pattern.split(text)):
        bold = part.startswith("**") and part.endswith("**")
        italic = part.startswith("*") and part.endswith("*") and not bold
        code = part.startswith("`") and part.endswith("`")
        clean = part[2:-2] if bold else part[1:-1] if italic or code else part
        run = paragraph.add_run(clean)
        set_run_font(run, base_size, bold=bold, italic=italic, color=base_color)
        if code:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            run.font.size = Pt(max(11, base_size - 2))


def add_figure(document, filename, caption):
    image_path = ASSETS / filename
    if not image_path.exists():
        raise FileNotFoundError(f"Missing screenshot: {image_path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    picture_width = 6.5 if filename == "program-workflow.png" else 6.85
    run.add_picture(str(image_path), width=Inches(picture_width))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    cap_run = cap.add_run(caption)
    set_run_font(cap_run, 12, italic=True, color=MUTED)


def configure_document(document, role_label):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(16)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05

    for style_name, size, color in (
        ("Title", 30, BLUE),
        ("Heading 1", 24, BLUE),
        ("Heading 2", 21, DARK),
        ("Heading 3", 18, DARK),
        ("Heading 4", 17, DARK),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run(f"ระบบขอแก้ไขข้อมูลออนไลน์  |  {role_label}")
    set_run_font(header_run, 11, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_cover(document, title, role_label):
    document.add_paragraph().paragraph_format.space_after = Pt(30)
    logo = document.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.add_run().add_picture(str(ROOT / "public" / "tsmlogo.png"), width=Inches(3.0))
    document.add_paragraph().paragraph_format.space_after = Pt(24)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run(title)
    set_run_font(run, 30, bold=True, color=BLUE)

    role = document.add_paragraph()
    role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    role.paragraph_format.space_after = Pt(20)
    role_run = role.add_run(role_label)
    set_run_font(role_run, 22, bold=True, color=DARK)

    byline = document.add_paragraph()
    byline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    byline_run = byline.add_run("จัดทำโดย แผนก IT Application Support")
    set_run_font(byline_run, 16, color=MUTED)
    document.add_page_break()


def add_table(document, rows):
    width = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(rows[0]):
        cell = header.cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, BLUE)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        set_run_font(run, 14, bold=True, color="FFFFFF")
    for source_row in rows[1:]:
        cells = table.add_row().cells
        for index in range(width):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if len(table.rows) % 2 == 1:
                shade_cell(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, source_row[index] if index < len(source_row) else "", 13)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def markdown_to_docx(source_path, output_path, role_label, images):
    lines = source_path.read_text(encoding="utf-8").splitlines()
    title = next((line[2:].strip() for line in lines if line.startswith("# ")), source_path.stem)
    document = Document()
    configure_document(document, role_label)
    add_cover(document, title, role_label)

    first_h1_skipped = False
    index = 0
    last_was_page_break = True
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue

        if stripped == "---":
            if not last_was_page_break:
                document.add_page_break()
                last_was_page_break = True
            index += 1
            continue

        if stripped.startswith("# ") and not first_h1_skipped:
            first_h1_skipped = True
            index += 1
            continue

        if stripped in ("**สำหรับ ผู้ขอ (Requester)**", "**สำหรับ ผู้อนุมัติ (Approver)**"):
            index += 1
            continue
        if stripped == "โดย แผนก IT Application Support":
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[index + 1]):
            table_lines = [stripped]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = []
            for table_line in table_lines:
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                rows.append([re.sub(r"[*_`]", "", cell) for cell in cells])
            add_table(document, rows)
            last_was_page_break = False
            continue

        caption_marker = stripped in images and stripped.startswith("*ภาพที่")
        if not caption_marker:
            heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
            if heading:
                level = len(heading.group(1))
                if level == 2 and not last_was_page_break:
                    document.add_page_break()
                paragraph = document.add_heading(level=max(1, level - 1))
                add_inline(paragraph, heading.group(2), {1: 24, 2: 21, 3: 18, 4: 17}[level])
            elif re.match(r"^-\s+", stripped):
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Inches(0.28)
                add_inline(paragraph, re.sub(r"^-\s+", "", stripped))
            elif re.match(r"^\d+\.\s+", stripped):
                paragraph = document.add_paragraph(style="List Number")
                paragraph.paragraph_format.left_indent = Inches(0.28)
                add_inline(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            elif stripped.startswith(">"):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.right_indent = Inches(0.15)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(7)
                paragraph.style = document.styles["Normal"]
                p_pr = paragraph._p.get_or_add_pPr()
                borders = OxmlElement("w:pBdr")
                left = OxmlElement("w:left")
                left.set(qn("w:val"), "single")
                left.set(qn("w:sz"), "18")
                left.set(qn("w:space"), "8")
                left.set(qn("w:color"), BLUE)
                borders.append(left)
                p_pr.append(borders)
                add_inline(paragraph, stripped.lstrip("> "), 15, MUTED)
            else:
                paragraph = document.add_paragraph()
                add_inline(paragraph, stripped)
            last_was_page_break = False

        if stripped in images:
            for filename, caption in images[stripped]:
                add_figure(document, filename, caption)
            last_was_page_break = False

        index += 1

    document.core_properties.title = title
    document.core_properties.subject = role_label
    document.core_properties.author = "IT Application Support"
    document.save(output_path)
    print(output_path)


markdown_to_docx(
    DOCS / "คู่มือการใช้งาน-ผู้ขอ.md",
    OUTPUT / "คู่มือการใช้งาน-ผู้ขอ-พร้อมภาพ.docx",
    "สำหรับ ผู้ขอ (Requester)",
    REQUESTER_IMAGES,
)
markdown_to_docx(
    DOCS / "คู่มือการใช้งาน-ผู้อนุมัติ.md",
    OUTPUT / "คู่มือการใช้งาน-ผู้อนุมัติ-พร้อมภาพ.docx",
    "สำหรับ ผู้อนุมัติ (Approver)",
    APPROVER_IMAGES,
)
